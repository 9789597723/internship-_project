
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify
from openpyxl import load_workbook
from pathlib import Path
import os

app = Flask(__name__)
app.secret_key = os.environ.get(
    "FLASK_SECRET_KEY", "change-this-for-local-demo"
)

# Project folder
BASE_DIR = Path(__file__).resolve().parent

# Your actual Excel filename
EXCEL_PATH = BASE_DIR / "Internship_Candidates_100(2).xlsx"

# Excel worksheet name
SHEET_NAME = "Candidates"


def get_sheet():
    """Open the Excel workbook and return its worksheet."""
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_PATH}\n"
            "Place Internship_Candidates_100(2).xlsx "
            "in the same folder as app.py."
        )

    wb = load_workbook(EXCEL_PATH)

    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb.active

    return wb, ws


def get_candidates():
    """Read candidate records from Excel."""
    wb, ws = get_sheet()

    try:
        headers = [cell.value for cell in ws[1]]
        candidates = []

        for row_num in range(2, ws.max_row + 1):
            values = [
                ws.cell(row_num, col).value
                for col in range(1, len(headers) + 1)
            ]

            if not any(value is not None for value in values):
                continue

            item = dict(zip(headers, values))
            item["_row"] = row_num

            for key, value in list(item.items()):
                if hasattr(value, "strftime"):
                    # Keep dates suitable for display.
                    if key == "Interview Date":
                        item[key] = value.strftime("%Y-%m-%d")
                    else:
                        item[key] = value.strftime("%d-%m-%Y")
                elif value is None:
                    item[key] = ""
                else:
                    item[key] = str(value)

            candidates.append(item)

        return headers, candidates

    finally:
        wb.close()


@app.route("/")
def index():
    statuses = [
        "Applied",
        "Under Review",
        "Shortlisted",
        "Interview Scheduled",
        "Selected",
        "Rejected",
    ]

    try:
        headers, all_candidates = get_candidates()

        # Search candidates by name, email, project, or other fields.
        q = request.args.get("q", "").strip().lower()
        status_filter = request.args.get("status", "").strip()

        candidates = all_candidates

        if q:
            candidates = [
                candidate
                for candidate in candidates
                if q in " ".join(
                    str(value)
                    for key, value in candidate.items()
                    if not key.startswith("_")
                ).lower()
            ]

        if status_filter:
            candidates = [
                candidate
                for candidate in candidates
                if candidate.get("Selection Status", "Applied")
                == status_filter
            ]

        total = len(all_candidates)

    except Exception as exc:
        flash(str(exc), "error")
        candidates = []
        total = 0

    return render_template(
        "index.html",
        candidates=candidates,
        statuses=statuses,
        total=total,
    )


@app.post("/update/<int:row_num>")
def update_candidate(row_num):
    wb = None

    try:
        wb, ws = get_sheet()

        if row_num < 2 or row_num > ws.max_row:
            flash("Candidate row not found.", "error")
            return redirect(url_for("index"))

        # Find the correct Excel columns by header name.
        headers = {
            str(ws.cell(1, col).value).strip(): col
            for col in range(1, ws.max_column + 1)
            if ws.cell(1, col).value is not None
        }

        interview_col = headers.get("Interview Date")
        status_col = headers.get("Selection Status")
        explanation_col = headers.get("AI Project Explanation")

        if interview_col:
            interview_date = request.form.get(
                "interview_date", ""
            ).strip()

            ws.cell(row_num, interview_col).value = (
                interview_date if interview_date else None
            )

        if status_col:
            ws.cell(row_num, status_col).value = request.form.get(
                "status", "Applied"
            ).strip()

        if explanation_col:
            ws.cell(row_num, explanation_col).value = (
                request.form.get("ai_explanation", "").strip()
                or "Not generated"
            )

        wb.save(EXCEL_PATH)
        flash("Candidate details saved to Excel.", "success")

    except Exception as exc:
        flash(f"Could not save changes: {exc}", "error")

    finally:
        if wb is not None:
            wb.close()

    return redirect(url_for("index"))


@app.post("/explain/<int:row_num>")
def explain_project(row_num):
    """Demo endpoint; a real AI API is not connected yet."""
    try:
        _, candidates = get_candidates()

        candidate = next(
            (
                item for item in candidates
                if item["_row"] == row_num
            ),
            None,
        )

        if not candidate:
            return jsonify({"error": "Candidate not found."}), 404

        title = candidate.get("Project Title", "") or "Untitled project"
        link = candidate.get("GitHub Link", "")

        if not link:
            explanation = (
                f"Project: {title}. GitHub link is missing. "
                "Add a valid repository link before reviewing this project."
            )
        else:
            explanation = (
                f"Project: {title}. Repository: {link}. "
                "This is a demo response, not an AI-generated project analysis. "
                "Connect an AI API and provide accessible project README/source "
                "content to generate a genuine explanation."
            )

        return jsonify({"explanation": explanation})

    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


if __name__ == "__main__":
    print("Excel file location:", EXCEL_PATH)
    print("Excel file exists:", EXCEL_PATH.exists())

    app.run(debug=True)
